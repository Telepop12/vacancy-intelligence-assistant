"""
Executive CV (English) → DOCX, опционально → PDF через установленный Word.

Содержимое живёт вне репозитория: data/resume/cv_en.yaml (личные данные).
Здесь только вёрстка: две страницы A4, шапка с контактами и опциональным фото,
секции с подчёркнутыми заголовками, настоящие списки Word.

    py -3.14 -m agents.cv_export                 # без фото и с фото, DOCX + PDF
    py -3.14 -m agents.cv_export --no-pdf        # только DOCX
    py -3.14 -m agents.cv_export --out output/cv # другая папка
"""
from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent.parent
CONTENT_PATH = BASE_DIR / "data" / "resume" / "cv_en.yaml"
DEFAULT_OUT = BASE_DIR / "data" / "resume" / "published" / "CV_EN"

ACCENT = RGBColor(0x1B, 0x2B, 0x5B)   # тёмно-синий, как в материалах лаборатории
MUTED = RGBColor(0x5C, 0x62, 0x70)
FONT = "Calibri"
BODY_PT = 9.5
SIDE_MARGIN_CM = 1.6
TEXT_WIDTH_CM = 21.0 - 2 * SIDE_MARGIN_CM


def _set_cell_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _bottom_border(paragraph, color: str = "1B2B5B", size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _spacing(paragraph, before: float = 0, after: float = 0, line: float | None = None) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing = line


class CvBuilder:
    def __init__(self, data: dict, photo: Path | None):
        self.d = data
        self.photo = photo
        self.doc = Document()
        self._setup_page()

    # -- базовая настройка -------------------------------------------------
    def _setup_page(self) -> None:
        for s in self.doc.sections:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
            s.top_margin = s.bottom_margin = Cm(1.4)
            s.left_margin = s.right_margin = Cm(SIDE_MARGIN_CM)
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(BODY_PT)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.0   # у Word по умолчанию 1,08
        # пустой абзац из шаблона убираем
        for p in list(self.doc.paragraphs):
            p._p.getparent().remove(p._p)

    # -- примитивы ---------------------------------------------------------
    def para(self, text: str = "", *, bold=False, italic=False, size=None,
             color: RGBColor | None = None, after=2.0, before=0.0, align=None):
        p = self.doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
        _spacing(p, before, after)
        if align is not None:
            p.alignment = align
        return p

    def heading(self, text: str) -> None:
        p = self.para(text.upper(), bold=True, size=10, color=ACCENT, before=6, after=2)
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            r.font.all_caps = True
        _bottom_border(p)

    def bullet(self, text: str, after: float = 1.0) -> None:
        p = self.doc.add_paragraph(style="List Bullet")
        p.add_run(text)
        fmt = p.paragraph_format
        fmt.left_indent = Cm(0.45)
        fmt.first_line_indent = Cm(-0.45)
        _spacing(p, 0, after)

    def role_line(self, company: str, location: str, dates: str) -> None:
        p = self.doc.add_paragraph()
        r = p.add_run(company.upper())
        r.bold = True
        r.font.color.rgb = ACCENT
        if location:
            r2 = p.add_run(f"  ·  {location}")
            r2.font.color.rgb = MUTED
        p.add_run("\t")
        r3 = p.add_run(dates)
        r3.bold = True
        p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.keep_with_next = True
        _spacing(p, 5, 0)

    # -- шапка -------------------------------------------------------------
    def header(self) -> None:
        d = self.d
        if self.photo and self.photo.exists():
            table = self.doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            _set_cell_borders_none(table)
            left, right = table.rows[0].cells
            left.width, right.width = Cm(13.9), Cm(3.5)
            self._header_text(left)
            rp = right.paragraphs[0]
            rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rp.add_run().add_picture(str(self.photo), width=Cm(2.7))
            _spacing(rp, 0, 0)
        else:
            self._header_text(None)

    def _header_text(self, cell) -> None:
        d = self.d
        # в ячейке таблицы уже есть один пустой абзац — используем его первым
        pending = [cell.paragraphs[0]] if cell is not None else []

        def add(text, *, bold=False, size=None, color=None, after=1.0):
            if pending:
                p = pending.pop()
            elif cell is not None:
                p = cell.add_paragraph()
            else:
                p = self.doc.add_paragraph()
            r = p.add_run(text)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
            _spacing(p, 0, after)
            return p

        add(d["name"], bold=True, size=20, color=ACCENT, after=1)
        add(d["title"], size=11.5, color=ACCENT, after=3)
        add(d["location"], size=9.5, color=MUTED, after=0)
        add("  ·  ".join(d["contacts"]), size=9.5, color=MUTED, after=4)

    # -- секции ------------------------------------------------------------
    def build(self) -> Document:
        d = self.d
        self.header()

        self.heading("Executive summary")
        self.para(d["summary"], after=2)

        self.heading("Career highlights")
        for h in d["highlights"]:
            self.bullet(h)

        self.heading("Core competencies")
        self.para("  ·  ".join(d["competencies"]), after=2)

        self.heading("Professional experience")
        for job in d["experience"]:
            self.role_line(job["company"], job.get("location", ""), job["dates"])
            self.para(job["role"], bold=True, after=1)
            if job.get("context"):
                self.para(job["context"], italic=True, color=MUTED, after=2)
            for b in job.get("bullets", []):
                self.bullet(b)
            if job.get("text"):
                self.para(job["text"], after=2)

        self.heading("Earlier career")
        for e in d["earlier"]:
            self.bullet(e)

        self.heading("Education & certifications")
        for e in d["education"]:
            self.para(e, after=1)
        for c in d["certifications"]:
            self.para(c, color=MUTED, size=9.5, after=0.5)

        self.heading("Languages & publications")
        self.para(d["languages"], after=1)
        self.para(d["publications"], after=0)
        return self.doc


def build_docx(out_path: Path, with_photo: bool) -> Path:
    data = yaml.safe_load(CONTENT_PATH.read_text(encoding="utf-8"))
    photo = (BASE_DIR / data["photo"]) if with_photo and data.get("photo") else None
    doc = CvBuilder(data, photo).build()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def docx_to_pdf(docx_path: Path) -> Path:
    """Конвертация через установленный Word (COM); LibreOffice на машине нет."""
    pdf_path = docx_path.with_suffix(".pdf")
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx_path}'); "
        f"$d.SaveAs2('{pdf_path}', 17); $d.Close(); $w.Quit(); "
        f"(Get-Item '{pdf_path}').Length"
    )
    res = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                         capture_output=True, text=True, timeout=120)
    if res.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(f"Word → PDF не удалась: {res.stderr.strip()[:300]}")
    return pdf_path


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--basename", default="Grishkin_CV_EN")
    args = ap.parse_args(argv)

    made: list[Path] = []
    for with_photo, suffix in ((False, ""), (True, "_photo")):
        docx_path = build_docx(args.out / f"{args.basename}{suffix}.docx", with_photo)
        made.append(docx_path)
        if not args.no_pdf:
            made.append(docx_to_pdf(docx_path))
    for p in made:
        print(f"{p}  ({p.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
