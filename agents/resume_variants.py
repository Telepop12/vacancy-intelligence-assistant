"""
Resume Variants Layer — стационарные варианты резюме под архетипы ролей.

hh.ru constraint: you respond to vacancies with a resume that ALREADY exists
in your account (a fresh upload sits in moderation ~10 minutes), so per-vacancy
resume generation is useless there. Instead the system maintains content for
2–5 permanent resume variants (config/resume_variants.yaml), each re-accenting
the SAME facts from the canonical resume for a different role archetype.

Hard rule enforced in the prompt: no invented facts. The LLM may reorder,
rephrase and emphasize — never add companies, numbers, projects or skills
that are not in the canonical resume.

Artifacts: output/resume_variants/{key}.json (state) + {key}.md (copy-paste
ready for hh.ru).
"""
from __future__ import annotations

import json
import os
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path

import yaml

BASE_DIR = Path(__file__).resolve().parent.parent
CONFIG_PATH = BASE_DIR / "config" / "resume_variants.yaml"
OUTPUT_DIR = BASE_DIR / "output" / "resume_variants"
FACTS_PATH = BASE_DIR / "data" / "resume" / "facts.md"


def load_facts() -> str:
    """Extra verified facts beyond the canonical resume (публикации, pet-проекты)."""
    if FACTS_PATH.exists():
        return FACTS_PATH.read_text(encoding="utf-8")
    return ""


@dataclass
class PositionAccent:
    company: str = ""
    period: str = ""
    title: str = ""
    bullets: list[str] = field(default_factory=list)


@dataclass
class ResumeVariant:
    key: str
    name: str = ""
    hh_title: str = ""
    about: str = ""
    key_skills: list[str] = field(default_factory=list)
    positions: list[PositionAccent] = field(default_factory=list)
    what_changed: str = ""          # чем этот вариант отличается от канона
    honest_gaps: str = ""           # чего в опыте объективно не хватает для архетипа
    generated_at: str = ""
    error: str = ""


# ---------------------------------------------------------------------------
# Config / state
# ---------------------------------------------------------------------------

def load_variants_config() -> dict:
    return yaml.safe_load(CONFIG_PATH.read_text(encoding="utf-8"))


def resume_path(cfg: dict | None = None) -> Path:
    cfg = cfg or load_variants_config()
    p = Path(cfg["resume_file"])
    return p if p.is_absolute() else BASE_DIR / p


def load_canonical_resume(cfg: dict | None = None) -> str:
    """Read the canonical resume text via the resume intelligence loader."""
    from agents.resume_intelligence import load_resume
    text, _src = load_resume(resume_path(cfg))
    return text


def load_saved_variants() -> dict[str, ResumeVariant]:
    """Read previously generated variants from output/resume_variants/."""
    variants: dict[str, ResumeVariant] = {}
    if not OUTPUT_DIR.exists():
        return variants
    for path in sorted(OUTPUT_DIR.glob("*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            data["positions"] = [PositionAccent(**p) for p in data.get("positions", [])]
            variants[data["key"]] = ResumeVariant(**data)
        except (json.JSONDecodeError, TypeError, KeyError):
            continue
    return variants


def _save_variant(variant: ResumeVariant) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / f"{variant.key}.json").write_text(
        json.dumps(asdict(variant), ensure_ascii=False, indent=1), encoding="utf-8"
    )
    (OUTPUT_DIR / f"{variant.key}.md").write_text(
        render_variant_md(variant), encoding="utf-8"
    )


def render_variant_md(v: ResumeVariant) -> str:
    """Human-readable, copy-paste-ready document for one variant."""
    lines = [
        f"# {v.name}",
        "",
        f"_Сгенерировано: {v.generated_at} · Источник фактов: каноническое резюме_",
        "",
        "## Желаемая должность (заголовок резюме на hh)",
        "",
        v.hh_title,
        "",
        "## Блок «О себе»",
        "",
        v.about,
        "",
        "## Ключевые навыки",
        "",
    ]
    lines += [f"- {s}" for s in v.key_skills]
    lines += ["", "## Акценты по позициям (что писать в описании опыта)"]
    for p in v.positions:
        lines += ["", f"### {p.company} · {p.period}", f"**{p.title}**", ""]
        lines += [f"- {b}" for b in p.bullets]
    lines += [
        "",
        "## Чем отличается от канонического резюме",
        "",
        v.what_changed,
        "",
        "## Честные пробелы под этот архетип",
        "",
        v.honest_gaps,
        "",
    ]
    return "\n".join(lines)


def export_variant_docx(variant: ResumeVariant, out_dir: Path | None = None) -> Path:
    """
    Export a variant to .docx formatted for hh.ru transfer.

    hh.ru's resume fields are plain text (no rich editor): Word lists, fonts
    and indents are lost on paste. So structure lives in the characters
    themselves — «— » dash bullets as plain paragraphs, blank lines between
    blocks, UPPERCASE section headers. Copy-paste keeps the visual structure.
    """
    from docx import Document
    from docx.shared import Pt

    out_dir = out_dir or OUTPUT_DIR / "docx"
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def para(text="", bold=False, size=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        return p

    para(v_title := (variant.hh_title or variant.name), bold=True, size=14)
    para(f"Вариант: {variant.name} · сгенерировано {variant.generated_at}", size=9)
    para()
    para("О СЕБЕ", bold=True)
    para(variant.about)
    para()
    para("КЛЮЧЕВЫЕ НАВЫКИ", bold=True)
    # hh key skills are entered one by one — keep them on one line for copying
    para(" · ".join(variant.key_skills))
    para()
    para("ОПЫТ РАБОТЫ (акценты по позициям)", bold=True)
    for pos in variant.positions:
        para()
        para(f"{pos.company} · {pos.period}", bold=True)
        para(pos.title)
        for b in pos.bullets:
            para(f"— {b}")

    path = out_dir / f"{variant.key}.docx"
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# LLM generation
# ---------------------------------------------------------------------------

_SYSTEM_PROMPT = """\
Отвечай исключительно на деловом русском языке, без канцелярита. Избегай английского
профессионального жаргона; допустимы общепринятые термины: AI, ERP, MLOps, ROI, P&L,
KPI, CEO, CIO, CTO, CDTO, CAIO, M&A, MBA, BI, MDM, ITSM, ITIL, DevOps, CI/CD.
Запрещены ИИ-штампы: «обладаю обширным опытом», «богатый опыт», «эффективный менеджер»,
«нацелен на результат» и подобные пустые формулы.
Запрещён контрастный оборот-финализация «X, а не Y» («…подтверждённая позиция,
а не декларация», «…результат, а не процесс», «не просто …, а …») — смысл пункта
должен быть ясен без такого противопоставления; просто не пиши вторую половину.

Ты — эксперт по executive-резюме для ролей CIO / CTO / CDTO / CAIO уровня Top Management.
Твоя задача — переакцентировать РЕАЛЬНОЕ резюме кандидата под заданный архетип роли.

ЖЕЛЕЗНОЕ ПРАВИЛО — НИКАКОЙ ВЫДУМКИ:
- используй ТОЛЬКО факты, компании, периоды, цифры и проекты из переданного резюме;
- нельзя добавлять навыки, технологии, достижения или масштабы, которых нет в тексте;
- цифры переноси точно как в резюме, не округляй и не «улучшай»;
- если под архетип фактов мало — честно скажи об этом в поле honest_gaps,
  а не заполняй пробел фантазией.

Что МОЖНО и НУЖНО делать:
- менять порядок и приоритет: релевантное архетипу — вперёд и подробнее;
- переформулировать операционный язык в executive-язык (результат и бизнес-эффект
  вместо перечня обязанностей), сохраняя фактическую точность;
- сжимать нерелевантное до 1–2 строк, не удаляя позиции целиком;
- подбирать формулировки с ключевыми словами архетипа, если они подтверждены опытом.

Отвечай строго валидным JSON без markdown и пояснений вокруг."""

_USER_TEMPLATE = """\
КАНОНИЧЕСКОЕ РЕЗЮМЕ (главный источник фактов, мастер-версия):
---
{resume_text}
---

ФАКТ-БАЗА (дополнительные подтверждённые факты: публикации, pet-проекты,
детализация позиций; внутри действует правило статусов — использовать можно
только ✅ и детализацию канонических позиций, разделы 🚫 игнорировать):
---
{facts}
---

АРХЕТИП ВАРИАНТА: {name}
Заголовок на hh (ориентир): {hh_title_hint}

ФОКУС (усилить, вывести вперёд):
{focus}

ПРИГЛУШИТЬ (сжать, но не выдумывать и не удалять позиции):
{deemphasize}

ТРЕБОВАНИЯ К РЕЗУЛЬТАТУ:
- about: блок «О себе» для hh, до {about_max_chars} символов, 3-5 плотных предложений,
  первое предложение — позиционирование под архетип;
- key_skills: ровно {key_skills_count} навыков, подтверждённых резюме,
  отсортированы по релевантности архетипу;
- positions: ВСЕ позиции из резюме в исходном хронологическом порядке (новые сверху);
  для каждой — название должности как в резюме (или уточнённое, если резюме допускает),
  и до {bullets_per_position} акцентов-буллетов под архетип; для старых/нерелевантных
  позиций достаточно 1-2 буллетов;
- what_changed: 2-4 предложения — какие акценты смещены относительно исходного резюме;
- honest_gaps: 1-3 предложения — чего в реальном опыте объективно не хватает
  для этого архетипа (это внутренняя заметка кандидату, не для публикации).

Верни ТОЛЬКО валидный JSON:
{{
  "hh_title": "заголовок резюме для hh",
  "about": "блок О себе",
  "key_skills": ["навык 1", "..."],
  "positions": [
    {{"company": "...", "period": "...", "title": "...", "bullets": ["...", "..."]}}
  ],
  "what_changed": "...",
  "honest_gaps": "..."
}}"""


def generate_variant(key: str, resume_text: str | None = None) -> ResumeVariant:
    """Generate (or regenerate) one variant. Saves artifacts on success."""
    cfg = load_variants_config()
    if key not in cfg["variants"]:
        return ResumeVariant(key=key, error=f"Неизвестный вариант: {key}")
    vcfg = cfg["variants"][key]

    variant = ResumeVariant(
        key=key,
        name=vcfg["name"],
        generated_at=datetime.now().isoformat(timespec="seconds"),
    )

    try:
        text = resume_text or load_canonical_resume(cfg)
    except Exception as exc:
        variant.error = f"Не удалось прочитать каноническое резюме: {exc}"
        return variant

    api_key = _get_api_key()
    if not api_key:
        variant.error = "Нет ANTHROPIC_API_KEY — генерация вариантов требует LLM"
        return variant

    prompt = _USER_TEMPLATE.format(
        resume_text=text,
        facts=load_facts() or "— факт-база пуста —",
        name=vcfg["name"],
        hh_title_hint=vcfg.get("hh_title_hint", ""),
        focus=vcfg.get("focus", ""),
        deemphasize=vcfg.get("deemphasize", ""),
        about_max_chars=cfg.get("about_max_chars", 900),
        key_skills_count=cfg.get("key_skills_count", 15),
        bullets_per_position=cfg.get("bullets_per_position", 4),
    )

    try:
        data = _call_llm(api_key, prompt)
        if _has_contrast_tic(data):
            # Deterministic guard: the model slips into the banned "X, а не Y"
            # rhetoric despite the system prompt — one retry with direct feedback.
            data = _call_llm(
                api_key,
                prompt + "\n\nВАЖНО: в прошлой попытке ты использовал запрещённый "
                "оборот «…, а не …». Перепиши все формулировки БЕЗ этого "
                "противопоставления — просто утверждение, без «а не»-части.",
            )
        variant.hh_title = data.get("hh_title", "")
        variant.about = data.get("about", "")
        variant.key_skills = data.get("key_skills") or []
        variant.positions = [
            PositionAccent(
                company=p.get("company", ""),
                period=p.get("period", ""),
                title=p.get("title", ""),
                bullets=p.get("bullets") or [],
            )
            for p in data.get("positions") or []
        ]
        variant.what_changed = data.get("what_changed", "")
        variant.honest_gaps = data.get("honest_gaps", "")
        _save_variant(variant)
    except Exception as exc:
        variant.error = f"Ошибка генерации: {exc}"

    return variant


def generate_all(resume_text: str | None = None) -> dict[str, ResumeVariant]:
    """Generate every configured variant sequentially. ~30-60s per variant."""
    cfg = load_variants_config()
    text = resume_text or load_canonical_resume(cfg)
    return {key: generate_variant(key, resume_text=text) for key in cfg["variants"]}


def _has_contrast_tic(data: dict) -> bool:
    """True if the banned '…, а не …' rhetoric appears in PUBLISHED fields."""
    published = [data.get("hh_title", ""), data.get("about", "")]
    published += data.get("key_skills") or []
    for p in data.get("positions") or []:
        published += p.get("bullets") or []
    return any(", а не " in s for s in published if isinstance(s, str))


def _call_llm(api_key: str, prompt: str) -> dict:
    import anthropic
    import httpx

    client = anthropic.Anthropic(api_key=api_key, http_client=httpx.Client(trust_env=False))
    response = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=8000,
        system=_SYSTEM_PROMPT,
        messages=[{"role": "user", "content": prompt}],
    )
    raw = response.content[0].text.strip()
    if raw.startswith("```"):
        lines = raw.split("\n")
        raw = "\n".join(lines[1:])
        if raw.endswith("```"):
            raw = raw[: raw.rfind("```")].rstrip()
    start, end = raw.find("{"), raw.rfind("}") + 1
    if start >= 0 and end > start:
        raw = raw[start:end]
    return json.loads(raw)


def _get_api_key() -> str:
    env_path = BASE_DIR / ".env"
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, _, v = line.partition("=")
                os.environ.setdefault(k.strip(), v.strip())
    key = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    return key if key and not key.startswith("sk-ant-ВАШ") else ""
